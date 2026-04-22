"""Excel I/O and column auto-detection for the warranty pipeline.

Responsibilities
----------------
1. Load the user's warranty `.xlsx` into a `pandas.DataFrame`.
2. Auto-detect the semantic role of each column using `config/column_aliases.yaml`
   (e.g., which column holds VIN, which holds tech_comment, ...). If detection
   fails, callers can fall back to an `ipywidgets` dropdown (handled in the
   notebook, not here).
3. Append analysis columns to the original workbook *without* destroying the
   user's formatting. We use `openpyxl.load_workbook` for this, so formulas,
   conditional formatting, and merged cells are preserved.

The column-mapping behavior is centralised here so every downstream step gets
consistent column names via a `ColumnMap` dataclass.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from . import failure_log

# ---------------------------------------------------------------------------
# Configuration loading
# ---------------------------------------------------------------------------


def _repo_root() -> Path:
    """Locate the repo root by walking up for a `config/` folder."""
    here = Path(__file__).resolve()
    for parent in [here, *here.parents]:
        if (parent / "config").is_dir() and (parent / "src").is_dir():
            return parent
    return Path.cwd()


def load_column_aliases(path: Optional[Path] = None) -> dict[str, list[str]]:
    """Load the English column-alias map used for auto-detection.

    Returns a dict like `{"vin": ["vin", "vin_no", ...], ...}`.
    """
    path = path or (_repo_root() / "config" / "column_aliases.yaml")
    with path.open("r", encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    # Normalize alias values to lower-case strings.
    return {role: [str(a).lower() for a in aliases] for role, aliases in data.items()}


def load_settings(path: Optional[Path] = None) -> dict[str, Any]:
    """Load the runtime settings YAML (alpha, model names, thresholds)."""
    path = path or (_repo_root() / "config" / "settings.yaml")
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


# ---------------------------------------------------------------------------
# Data containers
# ---------------------------------------------------------------------------


@dataclass
class ColumnMap:
    """Maps semantic roles to actual column names found in the workbook.

    Every attribute is `None` if no matching column was detected. Downstream
    modules should treat `None` as "feature unavailable" and degrade
    gracefully instead of raising.
    """

    vin: Optional[str] = None
    claim_id: Optional[str] = None
    repair_date: Optional[str] = None
    in_service_date: Optional[str] = None
    make: Optional[str] = None
    model: Optional[str] = None
    model_year: Optional[str] = None
    trim: Optional[str] = None
    engine_type: Optional[str] = None
    transmission_type: Optional[str] = None
    drivetrain: Optional[str] = None
    plant_code: Optional[str] = None
    production_date: Optional[str] = None
    mileage_at_repair: Optional[str] = None
    region: Optional[str] = None
    dealer_id: Optional[str] = None
    customer_complaint: Optional[str] = None
    tech_comment: Optional[str] = None
    replaced_part_number: Optional[str] = None
    repair_cost: Optional[str] = None
    labor_hours: Optional[str] = None

    def as_dict(self) -> dict[str, Optional[str]]:
        return self.__dict__.copy()

    def unresolved_roles(self) -> list[str]:
        """Roles for which no matching column was found."""
        return [role for role, col in self.as_dict().items() if col is None]


@dataclass
class WarrantyDataset:
    """Bundle of the loaded dataframe + detected column map + source path."""

    df: pd.DataFrame
    column_map: ColumnMap
    source_path: Path
    sheet_name: str
    extra: dict[str, Any] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Column auto-detection
# ---------------------------------------------------------------------------


def _normalize(name: str) -> str:
    """Lower-case + strip non-alphanumeric so 'Repair Date' == 'repair_date'."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def detect_columns(df: pd.DataFrame, aliases: dict[str, list[str]]) -> ColumnMap:
    """Auto-detect semantic column roles from the dataframe headers.

    The alias YAML lists canonical English spellings per role. We normalise
    both sides (strip punctuation, lowercase) before matching so the user's
    header can be "Repair Date" or "REPAIR_DATE" etc.
    """
    # Build a normalised lookup from header -> actual column name.
    header_lookup = {_normalize(c): c for c in df.columns}

    cmap = ColumnMap()
    for role, alias_list in aliases.items():
        if not hasattr(cmap, role):
            continue  # silently ignore roles that are not in ColumnMap yet
        for alias in alias_list:
            norm = _normalize(alias)
            if norm in header_lookup:
                setattr(cmap, role, header_lookup[norm])
                break
    return cmap


# ---------------------------------------------------------------------------
# Load / save
# ---------------------------------------------------------------------------


def load_warranty_xlsx(
    path: Path | str,
    sheet_name: str | int | None = 0,
    aliases: Optional[dict[str, list[str]]] = None,
) -> WarrantyDataset:
    """Load the user's warranty Excel and auto-detect column roles.

    Args:
        path: Path to the `.xlsx` the user dropped in `data/input/`.
        sheet_name: Worksheet to read (default: first sheet).
        aliases: Optional override of the column-alias map; defaults to
            `config/column_aliases.yaml`.

    Raises:
        FileNotFoundError: if `path` does not exist.
    """
    path = Path(path)
    if not path.exists():
        failure_log.record(
            step="io_utils.load_warranty_xlsx",
            context=f"path={path}",
            exc=FileNotFoundError(str(path)),
            lesson="Check the file was dropped into data/input/ before running the notebook.",
        )
        raise FileNotFoundError(path)

    try:
        df = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
    except Exception as e:
        failure_log.record(
            step="io_utils.load_warranty_xlsx",
            context=f"path={path}, sheet={sheet_name}",
            exc=e,
            lesson="Ensure the file is a valid .xlsx; save-as from .xls/.csv if needed.",
        )
        raise

    # `pd.read_excel` returns a dict when sheet_name is None; we require a df.
    if isinstance(df, dict):
        first_key = next(iter(df))
        df = df[first_key]
        sheet_name = first_key

    aliases = aliases or load_column_aliases()
    cmap = detect_columns(df, aliases)

    return WarrantyDataset(
        df=df,
        column_map=cmap,
        source_path=path,
        sheet_name=str(sheet_name) if sheet_name is not None else "Sheet1",
    )


def save_with_timestamp(
    src_path: Path | str,
    new_columns: pd.DataFrame,
    output_dir: Optional[Path] = None,
) -> Path:
    """Copy the original workbook, append new columns, save with timestamp.

    Implementation notes
    --------------------
    - We use `openpyxl.load_workbook` on the original file, so conditional
      formatting / merged cells / formulas are preserved.
    - New columns are appended after the last used column of the first
      worksheet. No existing cells are modified.
    - The output filename is
      `data/output/<original-stem>_analyzed_YYYYMMDD_HHMMSS.xlsx`.

    Args:
        src_path: Path to the original warranty .xlsx.
        new_columns: DataFrame whose rows align 1:1 with the original data
            rows. Its columns will be appended.
        output_dir: Override destination folder (defaults to data/output/).

    Returns:
        Absolute path to the saved file.
    """
    src_path = Path(src_path)
    output_dir = output_dir or (_repo_root() / "data" / "output")
    output_dir.mkdir(parents=True, exist_ok=True)

    ts = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = output_dir / f"{src_path.stem}_analyzed_{ts}.xlsx"

    try:
        wb = load_workbook(src_path)
        ws = wb.worksheets[0]  # First sheet — same as loaded by load_warranty_xlsx.

        # Determine where to start appending. openpyxl row 1 is the header.
        start_col = ws.max_column + 1
        header_row = 1
        data_start_row = 2

        # Write headers.
        for i, col_name in enumerate(new_columns.columns):
            ws.cell(row=header_row, column=start_col + i, value=str(col_name))

        # Write values. We align by positional index — the caller must ensure
        # new_columns has the same row order as the source sheet.
        for r_idx, (_, row) in enumerate(new_columns.iterrows()):
            for c_idx, col_name in enumerate(new_columns.columns):
                value = row[col_name]
                if pd.isna(value):
                    value = None
                ws.cell(
                    row=data_start_row + r_idx,
                    column=start_col + c_idx,
                    value=_coerce_xlsx_value(value),
                )

        wb.save(dest)
    except Exception as e:
        failure_log.record(
            step="io_utils.save_with_timestamp",
            context=f"src={src_path}",
            exc=e,
            lesson=(
                "Workbook append failed. Verify the source file is not open in Excel and "
                "that new_columns row count matches the source sheet row count."
            ),
        )
        raise

    return dest


def _coerce_xlsx_value(v: Any) -> Any:
    """Openpyxl accepts str/int/float/bool/None/datetime. Coerce lists/dicts."""
    if v is None or isinstance(v, (str, int, float, bool, _dt.datetime, _dt.date)):
        return v
    return str(v)


def write_summary_sheet(
    xlsx_path: Path,
    markdown: str,
    kpis: dict[str, Any],
    chart_paths: Optional[list[Path]] = None,
    sheet_name: str = "Summary",
) -> None:
    """Add a `Summary` sheet with briefing markdown, a KPI table, and charts.

    openpyxl cannot render markdown, so we write the text as wrapped cells
    and embed chart PNGs via `Image`. This is a best-effort presentation;
    the canonical briefing lives as `briefing_<ts>.md` / `.pdf`.
    """
    from openpyxl import load_workbook as _lw

    try:
        from openpyxl.drawing.image import Image
    except Exception:
        Image = None  # type: ignore

    wb = _lw(xlsx_path)
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    # KPI block.
    ws["A1"] = "KPI"
    ws["B1"] = "Value"
    for i, (k, v) in enumerate(kpis.items(), start=2):
        ws.cell(row=i, column=1, value=str(k))
        ws.cell(row=i, column=2, value=_coerce_xlsx_value(v))

    # Markdown block.
    md_start = len(kpis) + 3
    ws.cell(row=md_start, column=1, value="Briefing (markdown)")
    for j, line in enumerate(markdown.splitlines(), start=md_start + 1):
        ws.cell(row=j, column=1, value=line)

    # Widen column A for readability.
    ws.column_dimensions[get_column_letter(1)].width = 120

    # Embed chart images.
    if chart_paths and Image is not None:
        anchor_row = md_start + len(markdown.splitlines()) + 3
        for p in chart_paths:
            p = Path(p)
            if not p.exists():
                continue
            try:
                img = Image(str(p))
                ws.add_image(img, f"A{anchor_row}")
                anchor_row += 22  # ~22 rows per chart
            except Exception as e:  # pragma: no cover
                failure_log.record(
                    step="io_utils.write_summary_sheet",
                    context=f"chart={p}",
                    exc=e,
                    lesson="Chart embed failed; check the PNG is not corrupt.",
                )

    wb.save(xlsx_path)
